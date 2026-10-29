from pathlib import Path

from docx import Document as DocxDocument
from langchain_core.documents import Document as LangChainDocument
from pypdf import PdfReader


class UnstructuredParser:
    def parse(
        self,
        file_path: Path,
    ) -> list[LangChainDocument]:

        suffix = file_path.suffix.lower()

        if suffix == ".pdf":
            return self._parse_pdf(file_path)

        if suffix == ".docx":
            return self._parse_docx(file_path)

        if suffix == ".txt":
            return self._parse_txt(file_path)

        raise ValueError(
            f"Unsupported file extension: {suffix}"
        )

    def _build_document(
        self,
        text: str,
        file_path: Path,
        extra_metadata: dict | None = None,
    ) -> LangChainDocument | None:

        text = text.strip()

        if not text:
            return None

        metadata = {
            "source": str(file_path),
            "filename": file_path.name,
        }

        if extra_metadata:
            metadata.update(extra_metadata)

        return LangChainDocument(
            page_content=text,
            metadata=metadata,
        )

    def _parse_pdf(
        self,
        file_path: Path,
    ) -> list[LangChainDocument]:

        reader = PdfReader(str(file_path))
        documents = []

        for page_number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""

            document = self._build_document(
                text=text,
                file_path=file_path,
                extra_metadata={
                    "page_number": page_number,
                },
            )

            if document is not None:
                documents.append(document)

        return documents

    def _parse_docx(
        self,
        file_path: Path,
    ) -> list[LangChainDocument]:

        docx = DocxDocument(str(file_path))
        paragraphs = [
            paragraph.text.strip()
            for paragraph in docx.paragraphs
            if paragraph.text.strip()
        ]

        text = "\n".join(paragraphs)

        document = self._build_document(
            text=text,
            file_path=file_path,
        )

        return [document] if document is not None else []

    def _parse_txt(
        self,
        file_path: Path,
    ) -> list[LangChainDocument]:

        text = file_path.read_text(
            encoding="utf-8-sig",
            errors="replace",
        )

        document = self._build_document(
            text=text,
            file_path=file_path,
        )

        return [document] if document is not None else []
